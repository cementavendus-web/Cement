#!/usr/bin/env python3
"""Embed the rendered choropleth PNG as a 'World Map' sheet in the workbook,
so the map is visible on open with no live geocoding needed."""
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font

WORKBOOK = "SIPRI_Military_Expenditure_2025.xlsx"
IMG = "sipri_military_burden_map.png"

wb = openpyxl.load_workbook(WORKBOOK)
title = "World Map"
if title in wb.sheetnames:
    del wb[title]
ws = wb.create_sheet(title, index=0)          # first tab
ws.sheet_view.showGridLines = False

ws["A1"] = "World military burden, 2025 — military expenditure as a share of GDP"
ws["A1"].font = Font(bold=True, size=13, color="FF1C3F5F")
ws["A2"] = ("Source: SIPRI Military Expenditure Database, Apr. 2026. "
            "See the 'Military Burden Map' sheet for the underlying data.")
ws["A2"].font = Font(italic=True, size=10, color="FF5B6672")

img = XLImage(IMG)
# scale to ~1050 px wide, preserve aspect ratio
scale = 1050 / img.width
img.width = int(img.width * scale)
img.height = int(img.height * scale)
ws.add_image(img, "A4")

wb.save(WORKBOOK)
print(f"Embedded {IMG} ({img.width}x{img.height}) into '{title}' sheet.")
print("Sheets:", wb.sheetnames)
