#!/usr/bin/env python3
"""Add a map-ready 'Military Burden Map' sheet to the SIPRI workbook.

The sheet lists every country with its military burden (military expenditure
as a share of GDP), its SIPRI legend band, and a cell shaded in the matching
SIPRI colour. The Country + Burden columns are laid out so they can be fed
straight into Excel's Insert > Maps > Filled Map chart.
"""
import warnings

import geopandas as gpd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

WORKBOOK = "SIPRI_Military_Expenditure_2025.xlsx"

# --- Burden dataset (kept identical to build_map.py) -------------------------
BURDEN = {
    "ALB": 2.0, "AGO": 1.5, "ARG": 0.8, "ARM": 5.0, "AUT": 0.9, "AZE": 4.6,
    "BHS": 0.7, "BGD": 1.1, "BLR": 1.5, "BLZ": 1.2, "BEN": 0.7, "BOL": 1.4,
    "BIH": 0.8, "BWA": 2.9, "BRN": 3.0, "BGR": 1.9, "BFA": 5.0, "BDI": 2.0,
    "KHM": 2.2, "CMR": 1.2, "CAF": 1.5, "TCD": 2.5, "CHL": 1.9, "COG": 2.5,
    "HRV": 1.8, "CYP": 2.0, "CIV": 1.1, "COD": 1.1, "DJI": 3.5, "DOM": 0.7,
    "ECU": 2.4, "EGY": 1.2, "SLV": 1.0, "EST": 3.4, "ETH": 0.8, "FJI": 1.0,
    "GAB": 1.6, "GMB": 1.0, "GEO": 1.6, "GHA": 0.6, "GTM": 0.4, "GIN": 1.6,
    "GNB": 1.6, "GUY": 0.9, "HND": 1.7, "HUN": 2.1, "IRL": 0.2, "JAM": 1.4,
    "JOR": 4.7, "KAZ": 0.6, "KEN": 1.1, "KGZ": 1.4, "LVA": 3.2, "LBN": 3.0,
    "LSO": 1.8, "LBR": 0.6, "LTU": 2.9, "LUX": 0.7, "MDG": 0.6, "MWI": 0.8,
    "MYS": 1.0, "MLI": 4.0, "MRT": 2.4, "MDA": 0.6, "MNG": 0.7, "MNE": 1.6,
    "MAR": 4.3, "MOZ": 1.1, "MMR": 3.0, "NAM": 3.0, "NPL": 1.3, "NZL": 1.3,
    "NIC": 0.7, "NER": 2.5, "NGA": 0.9, "MKD": 1.7, "OMN": 5.6, "PAN": 0.0,
    "PNG": 0.5, "PRY": 1.0, "PER": 1.2, "PHL": 1.0, "PRT": 1.6, "RWA": 1.4,
    "SEN": 1.6, "SRB": 2.3, "SLE": 0.8, "SVK": 2.0, "SVN": 1.3, "ZAF": 0.7,
    "LKA": 1.6, "SUR": 0.6, "TZA": 1.0, "THA": 1.3, "TLS": 1.5, "TGO": 2.0,
    "TTO": 0.8, "TUN": 2.7, "UGA": 2.4, "URY": 1.9, "ZMB": 1.2, "ZWE": 0.9,
    "SWZ": 1.8, "CRI": 0.0, "ISL": 0.0, "SLB": 0.0, "VUT": 0.0,
}
TOP40_NAME_TO_ISO = {
    "United States": "USA", "China": "CHN", "Russia": "RUS", "Germany": "DEU",
    "India": "IND", "United Kingdom": "GBR", "Ukraine": "UKR",
    "Saudi Arabia": "SAU", "France": "FRA", "Japan": "JPN", "Israel": "ISR",
    "Italy": "ITA", "South Korea": "KOR", "Poland": "POL", "Spain": "ESP",
    "Canada": "CAN", "Australia": "AUS", "Turkiye": "TUR", "Netherlands": "NLD",
    "Algeria": "DZA", "Brazil": "BRA", "Taiwan": "TWN", "Singapore": "SGP",
    "Norway": "NOR", "Sweden": "SWE", "Indonesia": "IDN", "Denmark": "DNK",
    "Belgium": "BEL", "Colombia": "COL", "Mexico": "MEX", "Pakistan": "PAK",
    "Viet Nam": "VNM", "Romania": "ROU", "Greece": "GRC", "Kuwait": "KWT",
    "Finland": "FIN", "Switzerland": "CHE", "Iran": "IRN", "Czechia": "CZE",
    "Iraq": "IRQ",
}

# SIPRI legend: (label, low, high, ARGB fill). Blue = high burden, gold = low.
BINS = [
    ("≥ 9.0%",    9.0, float("inf"), "FF1C3F5F"),
    ("5.0%–9.0%", 5.0, 9.0,          "FF4F74A3"),
    ("4.0%–5.0%", 4.0, 5.0,          "FF6D93BD"),
    ("3.0%–4.0%", 3.0, 4.0,          "FF93B4D6"),
    ("2.0%–3.0%", 2.0, 3.0,          "FFC4D8E8"),
    ("1.0%–2.0%", 1.0, 2.0,          "FFDC9A34"),
    ("< 1%",      0.001, 1.0,        "FFECD992"),
    ("No spending", 0.0, 0.001,      "FFF5EFD8"),
]
NODATA_FILL = "FFD5D5D5"
# Dark fills need light text for legibility.
DARK_FILLS = {"FF1C3F5F", "FF4F74A3", "FF6D93BD", "FFDC9A34"}

# Natural Earth display-name fixes -> conventional country names.
NAME_FIX = {
    "United States of America": "United States",
    "Dem. Rep. Congo": "Dem. Rep. of the Congo",
    "Bosnia and Herz.": "Bosnia and Herzegovina",
    "Central African Rep.": "Central African Republic",
    "Dominican Rep.": "Dominican Republic",
    "Eq. Guinea": "Equatorial Guinea", "S. Sudan": "South Sudan",
    "W. Sahara": "Western Sahara", "N. Cyprus": "Northern Cyprus",
    "Solomon Is.": "Solomon Islands", "eSwatini": "Eswatini",
}
SKIP = {"ATA", "ATF", "FLK", "PRI", "NCL", "ATC"}


def band_for(v):
    if v is None:
        return None
    for i, (label, lo, hi, _c) in enumerate(BINS):
        if label == "No spending":
            if v == 0.0:
                return i
        elif lo <= v < hi:
            return i
    return None


def main():
    wb = openpyxl.load_workbook(WORKBOOK)
    ws_src = wb["Top 40 Spenders 2025"]
    for row in ws_src.iter_rows(min_row=2, values_only=True):
        if isinstance(row[0], int) and row[2] in TOP40_NAME_TO_ISO:
            BURDEN[TOP40_NAME_TO_ISO[row[2]]] = float(row[6])

    g = gpd.read_file(gpd.datasets.get_path("naturalearth_lowres"),
                      engine="pyogrio")
    rows = []
    for _, r in g.iterrows():
        iso, name = r["iso_a3"], r["name"]
        if iso in SKIP:
            continue
        name = NAME_FIX.get(name, name)
        v = BURDEN.get(iso)
        rows.append((name, iso if iso != "-99" else "", v))
    rows.sort(key=lambda t: t[0])

    # --- (re)build the sheet ---
    title = "Military Burden Map"
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title, index=1)

    thin = Side(style="thin", color="FFCED4DA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="FF1C3F5F")
    hdr_font = Font(bold=True, color="FFFFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    ws["A1"] = "World military burden, 2025 — military expenditure as a share of GDP (%)"
    ws["A1"].font = Font(bold=True, size=13, color="FF1C3F5F")
    ws["A2"] = ("Country + 'Burden (% of GDP)' feed Excel's Insert → Maps "
                "→ Filled Map. Cells are shaded to the SIPRI legend.")
    ws["A2"].font = Font(italic=True, size=10, color="FF5B6672")

    headers = ["Country", "ISO3", "Burden (% of GDP)", "Legend band"]
    hr = 4
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = center if c != 1 else left
        cell.border = border

    covered = 0
    for i, (name, iso, v) in enumerate(rows):
        rr = hr + 1 + i
        bi = band_for(v)
        if bi is None:
            band_label, fill = "No data", NODATA_FILL
        else:
            band_label, fill = BINS[bi][0], BINS[bi][3]
            covered += 1
        val = None if v is None else v

        ws.cell(row=rr, column=1, value=name).alignment = left
        ws.cell(row=rr, column=2, value=iso).alignment = center
        c3 = ws.cell(row=rr, column=3,
                     value=("" if val is None else round(val, 1)))
        c3.alignment = center
        c3.number_format = "0.0"
        c4 = ws.cell(row=rr, column=4, value=band_label)
        c4.alignment = center
        c4.fill = PatternFill("solid", fgColor=fill)
        c4.font = Font(color="FFFFFFFF" if fill in DARK_FILLS else "FF1B2430",
                       bold=(band_label == "≥ 9.0%"))
        for c in range(1, 5):
            ws.cell(row=rr, column=c).border = border

    last = hr + len(rows)

    # --- legend block to the right ---
    lg_col = 6
    ws.cell(row=hr, column=lg_col, value="Legend").font = Font(
        bold=True, color="FFFFFFFF")
    ws.cell(row=hr, column=lg_col).fill = hdr_fill
    ws.cell(row=hr, column=lg_col).alignment = center
    ws.cell(row=hr, column=lg_col + 1).fill = hdr_fill
    ws.cell(row=hr, column=lg_col).border = border
    ws.cell(row=hr, column=lg_col + 1).border = border
    legend_rows = [(lbl, fill) for (lbl, _lo, _hi, fill) in BINS]
    legend_rows.append(("No data", NODATA_FILL))
    for j, (lbl, fill) in enumerate(legend_rows):
        rr = hr + 1 + j
        sw = ws.cell(row=rr, column=lg_col)
        sw.fill = PatternFill("solid", fgColor=fill)
        sw.border = border
        lc = ws.cell(row=rr, column=lg_col + 1, value=lbl)
        lc.alignment = left
        lc.border = border

    # notes
    nr = last + 2
    notes = [
        "'Military burden' is military expenditure as a share of GDP; it "
        "represents the relative economic cost of the military.",
        "Countries shown as 'No data' are areas where data is unavailable or "
        "not reported.",
        "The 40 largest spenders use 2025 GDP-share figures; other countries "
        "use the latest available SIPRI figure.",
        "Source: SIPRI Military Expenditure Database, Apr. 2026. "
        "Geometry names: Natural Earth.",
    ]
    for k, txt in enumerate(notes):
        cell = ws.cell(row=nr + k, column=1, value=txt)
        cell.font = Font(italic=True, size=9, color="FF5B6672")

    # widths / freeze
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions[get_column_letter(lg_col)].width = 10
    ws.column_dimensions[get_column_letter(lg_col + 1)].width = 14
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False

    wb.save(WORKBOOK)
    print(f"Sheet '{title}' written: {len(rows)} countries, "
          f"{covered} with a burden value.")


if __name__ == "__main__":
    main()
